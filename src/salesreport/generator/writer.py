"""Write simulated rows as messy-but-realistic daily XLSX exports."""

from __future__ import annotations

import logging
import shutil
from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import Workbook

from salesreport import rules
from salesreport.dates import days_between, fmt
from salesreport.generator.events import GeneratorConfig, Simulation
from salesreport.generator.master import MasterData
from salesreport.inputs.files import day_folder

log = logging.getLogger(__name__)

COLUMNS: dict[str, list[str]] = {
    rules.STORE: [
        "Contract Date",
        "MSISDN",
        "Contract Type",
        "Status",
        "Sales Point",
        "Agent Login",
        "Contract ID",
    ],
    rules.SELFSERVICE: [
        "Contract ID",
        "Contract Date",
        "MSISDN",
        "Contract Type",
        "Status",
        "Sales Point",
        "Agent Login",
    ],
    rules.EXPRESS: ["Flow ID", "Flow Date", "Phone", "Agent", "Sales Point", "Step"],
    rules.CLICK_COLLECT: ["Order ID", "Order Date", "Phone Number", "Pickup Shop", "Order Status"],
    rules.WEB: ["Order No", "Order Date", "Phone", "Order Type", "Sales Point", "Channel"],
    rules.TARIFF: ["Change Date", "MSISDN", "Change Type", "Product"],
}


def write_table(
    path: Path,
    columns: list[str],
    rows: list[dict],
    *,
    title_rows: list[list[object]] | None = None,
    info_sheet: bool = False,
    footer_after_blanks: int = 0,
    upper_header: bool = False,
) -> None:
    """Write one export file, optionally with the kinds of mess real exports have."""
    workbook = Workbook(write_only=True)
    if info_sheet:
        info = workbook.create_sheet("Info")
        info.append(["This file was exported automatically."])
        info.append(["Data is on the next sheet."])
    sheet = workbook.create_sheet("Data")
    for title in title_rows or []:
        sheet.append(title)
    sheet.append([c.upper() for c in columns] if upper_header else columns)
    for row in rows:
        sheet.append([row.get(c) for c in columns])
    if footer_after_blanks:
        for _ in range(footer_after_blanks):
            sheet.append([])
        sheet.append(["Rows exported:", len(rows)])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def file_name(cfg: GeneratorConfig, source: str, day: date) -> str:
    """Normal name, or one of the deliberately inconsistent names from the config."""
    for d, s, prefix in cfg.cyrillic_names:
        if d == day and s == source:
            return f"{prefix}_{fmt(day)}.xlsx"
    for d, s, text in cfg.unpadded_names:
        if d == day and s == source:
            return f"{source}_{text}.xlsx"
    return f"{source}_{fmt(day)}.xlsx"


def write_daily(cfg: GeneratorConfig, sim: Simulation, data_dir: Path) -> int:
    """Write every source for every day; returns the number of files written."""
    daily = data_dir / "daily"
    if daily.exists():
        shutil.rmtree(daily)
    rng = np.random.default_rng([cfg.seed, 7])
    written = 0
    for day in days_between(cfg.start, cfg.end):
        folder = day_folder(data_dir, day)
        for source in rules.ALL_SOURCES:
            rows = list(sim.rows.get((source, day), []))
            order = rng.permutation(len(rows))
            rows = [rows[i] for i in order]
            mess = {
                "title_rows": (
                    [
                        [f"Daily export - {source.replace('_', ' ')}"],
                        [f"Export date {fmt(day)}"],
                        [],
                    ][: int(rng.integers(1, 4))]
                    if rng.random() < 0.25
                    else None
                ),
                "info_sheet": bool(rng.random() < 0.10),
                "footer_after_blanks": 6 if rng.random() < 0.15 else 0,
                "upper_header": bool(rng.random() < 0.10),
            }
            if source == rules.STORE and day == cfg.reexport_day and rows:
                # The first export stopped early; a partial re-export covers the rest of
                # the day with some overlap. config/overrides.csv says to use both.
                cut, overlap = int(len(rows) * 0.6), int(len(rows) * 0.5)
                write_table(
                    folder / file_name(cfg, source, day), COLUMNS[source], rows[:cut], **mess
                )
                write_table(
                    folder / f"{source}_{fmt(day)}_reexport.xlsx", COLUMNS[source], rows[overlap:]
                )
                written += 2
                continue
            write_table(folder / file_name(cfg, source, day), COLUMNS[source], rows, **mess)
            written += 1
    log.info("wrote %d daily files to %s", written, daily)
    return written


def write_master(master: MasterData, data_dir: Path) -> None:
    """Shop and agent master files (with a title row above the header)."""
    folder = data_dir / "master"
    folder.mkdir(parents=True, exist_ok=True)
    shops = [
        {
            "Shop ID": s.shop_id,
            "Shop Name": s.name,
            "Region": s.region,
            "Sales Points": ", ".join(s.sales_points),
        }
        for s in master.shops
    ]
    write_table(
        folder / "shops.xlsx",
        ["Shop ID", "Shop Name", "Region", "Sales Points"],
        shops,
        title_rows=[["Shop master (synthetic)"], []],
    )
    agents = [
        {
            "Agent Login": a.login,
            "Display Name": a.name or None,
            "Home Shop": a.home_shop or None,
            "Channel": a.channel,
        }
        for a in master.agents
    ]
    write_table(
        folder / "agents.xlsx",
        ["Agent Login", "Display Name", "Home Shop", "Channel"],
        agents,
        title_rows=[["Agent master (synthetic)"]],
    )


def write_manual_count(path: Path, counts: dict[str, int], start: date, end: date) -> None:
    """The independent manual count used on the reconciliation sheet."""
    rows = [{"Agent Login": k, "Manual Renewals": v} for k, v in sorted(counts.items())]
    write_table(
        path,
        ["Agent Login", "Manual Renewals"],
        rows,
        title_rows=[[f"Independent manual count {fmt(start)}-{fmt(end)} (synthetic)"], []],
    )
