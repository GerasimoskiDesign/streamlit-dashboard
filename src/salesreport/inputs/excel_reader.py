"""Robust reading of daily Excel exports.

Exports are produced by different tools and occasionally by hand, so the reader does
not trust the layout:

* the data may sit on any sheet (some files start with an "Info" sheet);
* title rows may sit above the header (``Daily export - store contracts``);
* a footer such as ``Rows exported: 312`` may follow a run of blank rows.

The header is therefore found by *content*: the first row (within the first rows of
any sheet) that contains every required column name after text normalisation. Reading
stops after a run of blank rows, so footers never become data.
"""

from __future__ import annotations

import itertools
import logging
from collections.abc import Iterable, Sequence
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from salesreport.normalize import normalize_text

log = logging.getLogger(__name__)

HEADER_SCAN_ROWS = 15
BLANK_RUN_STOP = 5


class HeaderNotFoundError(ValueError):
    """No sheet in the file has a row containing all required columns."""


def _is_blank(row: Sequence[object]) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row)


def find_header(
    rows: Sequence[Sequence[object]], required: Iterable[str], scan_rows: int = HEADER_SCAN_ROWS
) -> int | None:
    """Return the index of the first row that contains all ``required`` column names."""
    needed = {normalize_text(c) for c in required}
    for index, row in enumerate(rows[:scan_rows]):
        names = {normalize_text(c) for c in row if c is not None}
        if needed <= names:
            return index
    return None


def rows_until_blank_run(
    rows: Iterable[Sequence[object]], blank_run_stop: int = BLANK_RUN_STOP
) -> list[Sequence[object]]:
    """Collect rows, skipping isolated blanks, and stop at ``blank_run_stop`` blanks in a row."""
    kept: list[Sequence[object]] = []
    blanks = 0
    for row in rows:
        if _is_blank(row):
            blanks += 1
            if blanks >= blank_run_stop:
                break
            continue
        blanks = 0
        kept.append(row)
    return kept


def read_table(
    path: Path,
    required: Iterable[str],
    *,
    scan_rows: int = HEADER_SCAN_ROWS,
    blank_run_stop: int = BLANK_RUN_STOP,
) -> pd.DataFrame:
    """Read the table that contains ``required`` columns from an ``.xlsx`` file.

    Returns a DataFrame whose column names are normalised (``Contract Date`` ->
    ``contract date``). Cell values are returned untouched; typing happens in the loader.
    """
    required = list(required)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            head: list[Sequence[object]] = []
            for row in iterator:
                head.append(row)
                if len(head) >= scan_rows:
                    break
            header_index = find_header(head, required, scan_rows)
            if header_index is None:
                continue
            header = [normalize_text(c) for c in head[header_index]]
            body = rows_until_blank_run(
                itertools.chain(head[header_index + 1 :], iterator), blank_run_stop=blank_run_stop
            )
            width = len(header)
            records = [list(r[:width]) + [None] * (width - len(r)) for r in body]
            frame = pd.DataFrame(records, columns=header, dtype=object)
            frame = frame.loc[:, [c for c in frame.columns if c]]
            if header_index:
                log.debug(
                    "%s: header found on row %d of sheet %r",
                    path.name,
                    header_index + 1,
                    sheet.title,
                )
            return frame
    finally:
        workbook.close()
    raise HeaderNotFoundError(f"{path.name}: no sheet has all required columns {required}")
