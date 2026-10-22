"""The Excel report.

Layout of the main sheet: one block per shop (shop row, then its agents sorted by
volume), grouped by region; then an "Other channels" section; then a grand total.
Store / self-service / flag columns hold the engine's counts; *totals and shares are
Excel formulas* so the sheet stays live if someone filters or adds a note row.

Shop rows are written from the engine's own per-shop aggregation, not as a SUM of the
agent rows. That is deliberate: ``verify`` can then prove that both views agree
instead of checking a formula against itself.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from salesreport import rules
from salesreport.dates import fmt
from salesreport.engine.v4 import V4Result

# column letter -> (header, source column in the aggregated frame or None for formulas)
COLUMNS: list[tuple[str, str, str | None]] = [
    ("A", "Shop / channel", None),
    ("B", "Sales point(s)", None),
    ("C", "Agent", None),
    ("D", "Name", None),
    ("E", "Store contracts", "store"),
    ("F", "Self-service", "selfservice"),
    ("G", "Web w/o contract (full variant)", "web_nocontract"),
    ("H", "Total renewals", None),
    ("I", "Self-service digital", "ss_digital"),
    ("J", "Express", "express"),
    ("K", "Click & collect", "cc"),
    ("L", "Web order", "web"),
    ("M", "Digital renewals", "digital"),
    ("N", "Digital share", None),
    ("O", "Digital share excl. C&C", None),
    ("P", "C&C only", "cc_only"),
    ("Q", "Row type", None),
]
VALUE_COLUMNS = [(letter, source) for letter, _, source in COLUMNS if source]
SUM_LETTERS = [letter for letter, _ in VALUE_COLUMNS]
HEADER_ROW = 4
MAIN_SHEET = "Report"

INK = "1F2937"
FILL_HEADER = PatternFill("solid", fgColor=INK)
FILL_SECTION = PatternFill("solid", fgColor="F3F4F6")
FILL_BLOCK = PatternFill("solid", fgColor="E5E7EB")
FILL_TOTAL = PatternFill("solid", fgColor="D1D5DB")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True, color=INK)
MUTED = Font(color="6B7280")
THIN = Side(style="thin", color="9CA3AF")
DOUBLE = Side(style="double", color=INK)


@dataclass
class ReportInputs:
    """Everything the workbook shows."""

    v4: V4Result
    controls: pd.DataFrame  # section, metric, value, expectation, status
    evolution: pd.DataFrame
    comparison: pd.DataFrame
    reconciliation: pd.DataFrame | None
    reconciliation_note: str


def _row_formulas(ws: Worksheet, r: int) -> None:
    ws[f"H{r}"] = f"=E{r}+F{r}+G{r}"
    ws[f"N{r}"] = f'=IF(H{r}=0,"",M{r}/H{r})'
    ws[f"O{r}"] = f'=IF(H{r}=0,"",(M{r}-P{r})/H{r})'


def _write_counts(ws: Worksheet, r: int, row: pd.Series) -> None:
    for letter, source in VALUE_COLUMNS:
        ws[f"{letter}{r}"] = int(row[source])
    _row_formulas(ws, r)


def _style_row(ws: Worksheet, r: int, fill: PatternFill | None, font: Font | None) -> None:
    for letter, _, _ in COLUMNS:
        cell = ws[f"{letter}{r}"]
        if fill:
            cell.fill = fill
        if font:
            cell.font = font


def _ordered_blocks(result: V4Result) -> tuple[dict[str, list[pd.Series]], list[pd.Series]]:
    blocks = result.block_rows()
    shops = blocks.loc[blocks["block_type"] == "shop"].sort_values(
        ["region", "renewals", "block_id"], ascending=[True, False, True]
    )
    by_region: dict[str, list[pd.Series]] = {}
    for _, row in shops.iterrows():
        by_region.setdefault(row["region"], []).append(row)
    channels = blocks.loc[blocks["block_type"] == "channel"].copy()
    order = {f"CH:{c}": i for i, c in enumerate(rules.CHANNEL_ORDER)}
    channels["_order"] = channels["block_id"].map(order).fillna(99)
    return by_region, [row for _, row in channels.sort_values("_order").iterrows()]


def write_main_sheet(ws: Worksheet, result: V4Result, title: str, subtitle: str) -> None:
    """The block layout with formulas for totals and shares."""
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws["A2"] = subtitle
    ws["A2"].font = MUTED
    for letter, header, _ in COLUMNS:
        cell = ws[f"{letter}{HEADER_ROW}"]
        cell.value = header
        cell.fill = FILL_HEADER
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[HEADER_ROW].height = 32

    agents = result.agent_rows()
    by_region, channels = _ordered_blocks(result)
    r = HEADER_ROW + 1
    block_rows: list[int] = []

    def write_block(block: pd.Series, label: str) -> None:
        nonlocal r
        ws[f"A{r}"] = label
        ws[f"B{r}"] = block["sales_points"]
        ws[f"Q{r}"] = "block"
        _write_counts(ws, r, block)
        _style_row(ws, r, FILL_BLOCK, BOLD)
        block_rows.append(r)
        r += 1
        members = agents.loc[agents["block_id"] == block["block_id"]].sort_values(
            ["renewals", "agent"], ascending=[False, True]
        )
        for _, agent in members.iterrows():
            ws[f"A{r}"] = label
            ws[f"A{r}"].font = MUTED
            ws[f"B{r}"] = agent["sales_points"]
            ws[f"C{r}"] = agent["agent"]
            ws[f"D{r}"] = agent["agent_name"]
            ws[f"Q{r}"] = "agent"
            _write_counts(ws, r, agent)
            r += 1

    def section(text: str) -> None:
        nonlocal r
        ws[f"A{r}"] = text
        ws[f"Q{r}"] = "section"
        _style_row(ws, r, FILL_SECTION, Font(bold=True, italic=True, color=INK))
        r += 1

    # Regions in master-data order (North, South, East, West), not alphabetically.
    for region in sorted(by_region, key=lambda name: min(b["block_id"] for b in by_region[name])):
        section(f"Region {region}")
        for block in by_region[region]:
            write_block(block, f"{block['block_name']} ({block['block_id']})")
    if channels:
        section("Other channels")
        for block in channels:
            write_block(block, str(block["block_name"]))

    ws[f"A{r}"] = "Grand total"
    ws[f"Q{r}"] = "total"
    for letter in SUM_LETTERS:
        refs = ",".join(f"{letter}{b}" for b in block_rows) or "0"
        ws[f"{letter}{r}"] = f"=SUM({refs})"
    _row_formulas(ws, r)
    _style_row(ws, r, FILL_TOTAL, BOLD)
    for letter, _, _ in COLUMNS:
        ws[f"{letter}{r}"].border = Border(top=DOUBLE)

    last = r
    for row in ws.iter_rows(min_row=HEADER_ROW + 1, max_row=last):
        for cell in row:
            if cell.column_letter in SUM_LETTERS or cell.column_letter == "H":
                cell.number_format = "#,##0"
            elif cell.column_letter in ("N", "O"):
                cell.number_format = "0.0%"
            elif cell.column_letter == "Q":
                cell.font = Font(color="9CA3AF", size=9)
    widths = {"A": 30, "B": 16, "C": 12, "D": 22, "Q": 9}
    for letter, _, _ in COLUMNS:
        ws.column_dimensions[letter].width = widths.get(letter, 12)
    if result.variant == rules.VARIANT_STRICT:
        ws.column_dimensions["G"].hidden = True
    ws.freeze_panes = f"E{HEADER_ROW + 1}"


def write_frame(ws: Worksheet, frame: pd.DataFrame, title: str, note: str = "") -> None:
    """Simple styled table sheet."""
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color=INK)
    if note:
        ws["A2"] = note
        ws["A2"].font = MUTED
    header_row = 4
    for c, name in enumerate(frame.columns, start=1):
        cell = ws.cell(row=header_row, column=c, value=str(name))
        cell.fill = FILL_HEADER
        cell.font = WHITE_BOLD
    for i, values in enumerate(frame.itertuples(index=False), start=header_row + 1):
        for c, value in enumerate(values, start=1):
            if isinstance(value, pd.Timestamp):
                value = value.to_pydatetime().date()
            elif value is not None and not isinstance(value, str | int | float | bool):
                value = value.item() if hasattr(value, "item") else str(value)
            if isinstance(value, float) and value != value:
                value = None
            cell = ws.cell(row=i, column=c, value=value)
            if hasattr(value, "year"):
                cell.number_format = "DD.MM.YYYY"
    for c, name in enumerate(frame.columns, start=1):
        sample = [len(str(name))] + [len(str(v)) for v in frame.iloc[:200, c - 1]]
        ws.column_dimensions[get_column_letter(c)].width = min(max(sample) + 2, 70)
    ws.freeze_panes = f"A{header_row + 1}"


def write_workbook(inputs: ReportInputs, path: Path) -> Path:
    """Write the full report workbook."""
    result = inputs.v4
    period = result.period
    workbook = Workbook()
    main = workbook.active
    main.title = MAIN_SHEET
    write_main_sheet(
        main,
        result,
        f"Renewals and digital share - month to date {fmt(period.start)} - {fmt(period.end)}",
        f"Logic v4 (locked) | variant: {result.variant} | "
        f"web look-ahead: {result.lookahead_days} days"
        " | synthetic data",
    )

    unmatched = result.unmatched.copy()
    unmatched["window_open"] = unmatched["window_open"].map({True: "yes", False: ""})
    write_frame(
        workbook.create_sheet("Unmatched digital"),
        unmatched.rename(columns={"window_open": "look-ahead window still open"}),
        "Digital transactions without a matching renewal",
        "Rule 6: nothing is dropped silently. Reason a = excluded population, "
        "b = contract exists but is not a renewal, c = no contract that day.",
    )
    write_frame(
        workbook.create_sheet("Control totals"),
        inputs.controls,
        "Control totals",
        "Rows with an expectation must show OK before numbers are published.",
    )
    write_frame(
        workbook.create_sheet("Web look-ahead links"),
        result.links,
        "Web orders linked to a renewal booked later (rule 7)",
        f"Window: {result.lookahead_days} days. before_period = order placed before "
        f"{fmt(period.start)}; the renewal still counts in this month.",
    )
    write_frame(
        workbook.create_sheet("v3 vs v4"),
        inputs.evolution,
        "Same inputs, four logic versions",
        "v3 = prototype at scale (each digital source counted separately and added). "
        "v4 = event-level engine. Per-shop detail starts below the version table.",
    )
    sheet = workbook["v3 vs v4"]
    start_row = len(inputs.evolution) + 8
    sheet.cell(row=start_row - 1, column=1, value="Per shop / channel").font = BOLD
    for c, name in enumerate(inputs.comparison.columns, start=1):
        cell = sheet.cell(row=start_row, column=c, value=str(name))
        cell.fill = FILL_HEADER
        cell.font = WHITE_BOLD
    for i, values in enumerate(inputs.comparison.itertuples(index=False), start=start_row + 1):
        for c, value in enumerate(values, start=1):
            if isinstance(value, float) and value != value:
                value = None
            sheet.cell(row=i, column=c, value=value.item() if hasattr(value, "item") else value)

    reconciliation = inputs.reconciliation
    write_frame(
        workbook.create_sheet("Reconciliation"),
        reconciliation
        if reconciliation is not None
        else pd.DataFrame({"note": ["no manual count found"]}),
        "Engine vs independent manual count",
        inputs.reconciliation_note,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
