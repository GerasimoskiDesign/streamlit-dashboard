"""``python -m salesreport verify <workbook>`` - check the published workbook itself.

The engine guarantees its invariants by construction, but people open, edit, filter
and forward the Excel file. Verification therefore reads the *workbook* (evaluating
its formulas) and checks the three things a reader relies on:

1. every shop/channel block row equals the sum of its agent rows (all count columns);
2. no row has more digital renewals than total renewals;
3. the grand total equals the sum of the block rows.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from salesreport.report.excel import COLUMNS, HEADER_ROW, MAIN_SHEET, SUM_LETTERS
from salesreport.report.formula import FormulaError, SheetEvaluator

CHECK_LETTERS = [*SUM_LETTERS, "H"]
HEADERS = {letter: header for letter, header, _ in COLUMNS}


@dataclass
class VerifyResult:
    """Outcome of a workbook verification."""

    failures: list[str] = field(default_factory=list)
    blocks_checked: int = 0
    rows_checked: int = 0

    @property
    def ok(self) -> bool:
        return not self.failures


def _close(a: float, b: float) -> bool:
    return abs(a - b) < 1e-9


def verify_workbook(path: Path) -> VerifyResult:
    """Run all checks on the main sheet of ``path``."""
    result = VerifyResult()
    workbook = load_workbook(path, data_only=False)
    try:
        if MAIN_SHEET not in workbook.sheetnames:
            result.failures.append(f"sheet {MAIN_SHEET!r} not found")
            return result
        ws = workbook[MAIN_SHEET]
        evaluator = SheetEvaluator(ws)

        def number(letter: str, row: int) -> float:
            value = evaluator.value(f"{letter}{row}")
            if value in (None, ""):
                return 0.0
            if isinstance(value, bool) or not isinstance(value, int | float):
                raise FormulaError(f"{letter}{row} is not a number: {value!r}")
            return float(value)

        def label(row: int) -> str:
            parts = [str(ws[f"A{row}"].value or "")]
            if ws[f"C{row}"].value:
                parts.append(str(ws[f"C{row}"].value))
            return " / ".join(parts)

        blocks: list[tuple[int, list[int]]] = []
        total_row: int | None = None
        for row in range(HEADER_ROW + 1, ws.max_row + 1):
            kind = ws[f"Q{row}"].value
            if kind == "block":
                blocks.append((row, []))
            elif kind == "agent":
                if not blocks:
                    result.failures.append(f"row {row}: agent row before any block row")
                    continue
                blocks[-1][1].append(row)
            elif kind == "total":
                total_row = row

        try:
            for block_row, agent_rows in blocks:
                result.blocks_checked += 1
                for letter in CHECK_LETTERS:
                    block_value = number(letter, block_row)
                    agent_sum = sum(number(letter, r) for r in agent_rows)
                    if not _close(block_value, agent_sum):
                        result.failures.append(
                            f"row {block_row} ({label(block_row)}): {HEADERS[letter]} block = "
                            f"{block_value:g} but agent rows sum to {agent_sum:g}"
                        )
            rows_to_check = [r for b, agents in blocks for r in (b, *agents)]
            if total_row is not None:
                rows_to_check.append(total_row)
            for row in rows_to_check:
                result.rows_checked += 1
                digital, total = number("M", row), number("H", row)
                if digital > total:
                    result.failures.append(
                        f"row {row} ({label(row)}): digital {digital:g} > total {total:g}"
                    )
            if total_row is None:
                result.failures.append("grand total row not found")
            else:
                for letter in CHECK_LETTERS:
                    grand = number(letter, total_row)
                    block_sum = sum(number(letter, b) for b, _ in blocks)
                    if not _close(grand, block_sum):
                        result.failures.append(
                            f"row {total_row} (Grand total): {HEADERS[letter]} = {grand:g} but "
                            f"blocks sum to {block_sum:g}"
                        )
        except FormulaError as exc:
            result.failures.append(f"could not evaluate workbook: {exc}")
    finally:
        workbook.close()
    return result
