"""Production hardening: logic lock, workbook verification, formula evaluation, CLI."""

from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from salesreport import cli
from salesreport.lock import (
    LOCKED_FILES,
    MANIFEST_NAME,
    LogicModifiedError,
    check_manifest,
    verify_lock,
    write_manifest,
)
from salesreport.report.formula import evaluate
from salesreport.report.verify import verify_workbook
from salesreport.settings import PROJECT_ROOT


# --- hash lock ----------------------------------------------------------------------------
def test_committed_manifest_matches_the_engine():
    assert check_manifest(PROJECT_ROOT) == []


@pytest.fixture
def locked_copy(tmp_path: Path) -> Path:
    for name in LOCKED_FILES:
        target = tmp_path / name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(PROJECT_ROOT / name, target)
    write_manifest(tmp_path)
    return tmp_path


def test_lock_refuses_a_modified_engine_file(locked_copy: Path):
    verify_lock(locked_copy)
    engine = locked_copy / "src/salesreport/engine/v4.py"
    engine.write_text(engine.read_text(encoding="utf-8") + "\n# tweak\n", encoding="utf-8")
    with pytest.raises(LogicModifiedError, match="logic was modified"):
        verify_lock(locked_copy)
    assert check_manifest(locked_copy) == ["changed: src/salesreport/engine/v4.py"]


def test_lock_ignores_line_ending_differences(locked_copy: Path):
    rules_file = locked_copy / "src/salesreport/rules.py"
    rules_file.write_bytes(rules_file.read_bytes().replace(b"\n", b"\r\n"))
    verify_lock(locked_copy)


def test_lock_after_review_accepts_the_change(locked_copy: Path):
    (locked_copy / "src/salesreport/rules.py").write_text("# reviewed change\n", encoding="utf-8")
    with pytest.raises(LogicModifiedError):
        verify_lock(locked_copy)
    write_manifest(locked_copy)
    verify_lock(locked_copy)


def test_missing_manifest_refuses(tmp_path: Path):
    with pytest.raises(LogicModifiedError, match=MANIFEST_NAME):
        verify_lock(tmp_path)


# --- formula evaluator ----------------------------------------------------------------------
def test_formula_subset():
    cells = {"A1": 2, "A2": 3, "B1": 0, "C1": None}
    get = cells.get
    assert evaluate("=A1+A2", get) == 5
    assert evaluate("=SUM(A1:A2)", get) == 5
    assert evaluate("=SUM(A1,A2,C1)", get) == 5
    assert evaluate('=IF(B1=0,"",A1/B1)', get) == ""
    assert evaluate('=IF(A2=0,"",(A2-A1)/A2)', get) == pytest.approx(1 / 3)


# --- verify ---------------------------------------------------------------------------------
def test_verify_accepts_the_generated_workbook(built):
    result = verify_workbook(built.excel_path)
    assert result.ok, result.failures
    assert result.blocks_checked >= 30


def _first_rows(ws, kind: str) -> list[int]:
    return [r for r in range(5, ws.max_row + 1) if ws[f"Q{r}"].value == kind]


def test_verify_catches_an_edited_agent_row(built, tmp_path: Path):
    broken = tmp_path / "broken.xlsx"
    workbook = load_workbook(built.excel_path)
    ws = workbook["Report"]
    agent_row = _first_rows(ws, "agent")[0]
    ws[f"E{agent_row}"] = ws[f"E{agent_row}"].value + 5
    workbook.save(broken)
    result = verify_workbook(broken)
    assert not result.ok
    assert any(f"row {agent_row - 1}" in f and "Store contracts" in f for f in result.failures)


def test_verify_catches_digital_above_total(built, tmp_path: Path):
    broken = tmp_path / "broken.xlsx"
    workbook = load_workbook(built.excel_path)
    ws = workbook["Report"]
    agent_row = _first_rows(ws, "agent")[0]
    ws[f"M{agent_row}"] = 10_000
    workbook.save(broken)
    result = verify_workbook(broken)
    assert any(f"row {agent_row} " in f and "digital 10000 > total" in f for f in result.failures)


def test_verify_catches_a_wrong_grand_total(built, tmp_path: Path):
    broken = tmp_path / "broken.xlsx"
    workbook = load_workbook(built.excel_path)
    ws = workbook["Report"]
    total_row = _first_rows(ws, "total")[0]
    ws[f"E{total_row}"] = 1
    workbook.save(broken)
    result = verify_workbook(broken)
    assert any("Grand total" in f and "Store contracts" in f for f in result.failures)


def test_workbook_uses_live_formulas_and_has_all_sheets(built):
    workbook = load_workbook(built.excel_path)
    assert workbook.sheetnames == [
        "Report",
        "Unmatched digital",
        "Control totals",
        "Web look-ahead links",
        "v3 vs v4",
        "Reconciliation",
    ]
    ws = workbook["Report"]
    row = _first_rows(ws, "block")[0]
    assert ws[f"H{row}"].value == f"=E{row}+F{row}+G{row}"
    assert str(ws[f"N{row}"].value).startswith("=IF(")
    total_row = _first_rows(ws, "total")[0]
    assert str(ws[f"E{total_row}"].value).startswith("=SUM(")


# --- CLI ------------------------------------------------------------------------------------
def test_cli_verify_returns_zero_for_a_good_workbook(built):
    assert cli.main(["verify", str(built.excel_path)]) == 0


def test_cli_build_refuses_incomplete_inputs(generated, tmp_path: Path):
    _, data_dir = generated
    sparse = tmp_path / "data"
    shutil.copytree(data_dir / "master", sparse / "master")
    shutil.copytree(
        data_dir / "daily" / "September" / "01.09.2026",
        sparse / "daily" / "September" / "01.09.2026",
    )
    code = cli.main(
        [
            "build",
            "03.09.2026",
            "--data-dir",
            str(sparse),
            "--out-dir",
            str(tmp_path / "out"),
            "--history",
            "0",
        ]
    )
    assert code == 4
    assert not (tmp_path / "out").exists()


def test_cli_rejects_a_bad_date(generated, tmp_path: Path):
    _, data_dir = generated
    assert (
        cli.main(["build", "31.02.2026", "--data-dir", str(data_dir), "--out-dir", str(tmp_path)])
        == 2
    )


def test_cli_full_variant_build(generated, tmp_path: Path):
    _, data_dir = generated
    code = cli.main(
        [
            "build",
            "17.09",
            "--variant",
            "full",
            "--history",
            "0",
            "--data-dir",
            str(data_dir),
            "--out-dir",
            str(tmp_path),
        ]
    )
    assert code == 0
    assert (tmp_path / f"sales_report_{date(2026, 9, 17):%d.%m.%Y}.xlsx").exists()
